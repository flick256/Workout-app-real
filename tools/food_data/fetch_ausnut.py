#!/usr/bin/env python3
"""
Builds app/src/main/assets/generic_foods.json from FSANZ's AUSNUT 2023 food nutrient
database (3,700+ Australian foods "as eaten": breads, takeaway, home-cooked dishes...).

AUSNUT 2023 © Food Standards Australia New Zealand, used under its licence (based on
CC BY-SA 3.0 AU): https://www.foodstandards.gov.au/science-data/food-nutrient-databases/ausnut

Run by .github/workflows/food-data.yml (this machine needs internet). It finds the
current "Food nutrient profiles" spreadsheet on FSANZ's data files page, so a new file
name on their side doesn't break it. Usage: python3 fetch_ausnut.py [out.json]
"""
import io
import json
import re
import sys
import urllib.request
from urllib.parse import urljoin

import openpyxl

PAGES = [
    "https://www.foodstandards.gov.au/science-data/food-nutrient-databases/ausnut/data-files",
    "https://www.foodstandards.gov.au/science-data/food-nutrient-databases/ausnut/food-nutrients",
]
UA = "Mozilla/5.0 (X11; Linux x86_64) Forge-food-data/1.0 (personal app; build script)"
KJ_PER_KCAL = 4.184


def get(url):
    req = urllib.request.Request(url, headers={"User-Agent": UA})
    with urllib.request.urlopen(req, timeout=60) as r:
        return r.read()


def all_links():
    links = []
    for page in PAGES:
        try:
            html = get(page).decode("utf-8", "replace")
        except Exception as e:  # noqa: BLE001
            print(f"!! couldn't read {page}: {e}")
            continue
        for href in re.findall(r'href="([^"]+\.xlsx[^"]*)"', html, flags=re.I):
            links.append(urljoin(page, href.replace("&amp;", "&")))
    links = list(dict.fromkeys(links))
    print("Spreadsheets found:")
    for link in links:
        print("  ", link)
    return links


def find_nutrient_file(links):
    def score(link):
        l = link.lower().replace("%20", " ")
        s = 0
        if "ausnut" in l: s += 2
        if "2023" in l: s += 2
        if "nutrient" in l: s += 3
        if "profile" in l or "per 100" in l: s += 3
        for bad in ("measure", "supplement", "recipe", "detail", "list of nutrients", "guideline", "classification"):
            if bad in l: s -= 5
        return s
    if not links:
        sys.exit("No spreadsheet links found on the FSANZ pages")
    best = max(links, key=score)
    print("Using:", best)
    return best


def pick(headers, *patterns, avoid=()):
    """Index of the first header matching all words of a pattern (tried in order)."""
    low = [(h or "").lower() for h in headers]
    for pattern in patterns:
        words = pattern.lower().split("|")
        for i, h in enumerate(low):
            if all(w in h for w in words) and not any(a in h for a in avoid):
                return i
    return None


def main():
    out = sys.argv[1] if len(sys.argv) > 1 else "app/src/main/assets/generic_foods.json"
    links = all_links()
    wb = openpyxl.load_workbook(io.BytesIO(get(find_nutrient_file(links))), read_only=True, data_only=True)
    print("Sheets:", wb.sheetnames)
    best = None
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        for r, row in enumerate(rows[:15]):
            cells = [str(c) if c is not None else "" for c in row]
            if pick(cells, "food name") is not None and pick(cells, "protein") is not None:
                best = (ws.title, r, cells, rows)
                break
        if best:
            break
    if not best:
        sys.exit("Couldn't find a header row with 'Food Name' and 'Protein'")
    title, r, headers, rows = best
    print(f"Sheet '{title}', header row {r + 1}:")
    for i, h in enumerate(headers):
        print(f"  [{i}] {h}")

    col = {
        "key": pick(headers, "public food key", "food key", "food id"),
        "name": pick(headers, "food name"),
        "kj": pick(headers, "energy|with dietary fibre|kj", "energy|including dietary fibre|kj", "energy|kj"),
        "protein": pick(headers, "protein"),
        "fat": pick(headers, "fat, total", "total fat", "fat|(g)", avoid=("acid", "trans", "saturated")),
        "carbs": pick(headers, "available carbohydrate, with sugar alcohols", "available carbohydrate", "carbohydrate"),
        "sugars": pick(headers, "total sugars", "sugars"),
        "fibre": pick(headers, "total dietary fibre", "dietary fibre", avoid=("energy",)),
        "sodium": pick(headers, "sodium"),
    }
    print("Columns used:", {k: (v, headers[v] if v is not None else None) for k, v in col.items()})
    for required in ("name", "kj", "protein", "fat", "carbs"):
        if col[required] is None:
            sys.exit(f"Missing column: {required}")

    def num(row, key):
        i = col[key]
        if i is None or i >= len(row):
            return None
        v = row[i]
        try:
            return float(v)
        except (TypeError, ValueError):
            return None

    foods = []
    for row in rows[r + 1:]:
        name = row[col["name"]] if col["name"] < len(row) else None
        kj = num(row, "kj")
        if not name or kj is None:
            continue
        food = {
            "k": str(row[col["key"]]).strip() if col["key"] is not None and row[col["key"]] else str(len(foods)),
            "n": " ".join(str(name).split()),
            "e": round(kj / KJ_PER_KCAL, 1),
            "p": round(num(row, "protein") or 0.0, 2),
            "c": round(num(row, "carbs") or 0.0, 2),
            "f": round(num(row, "fat") or 0.0, 2),
        }
        for short, key in (("fi", "fibre"), ("s", "sugars")):
            v = num(row, key)
            if v is not None:
                food[short] = round(v, 2)
        na = num(row, "sodium")
        if na is not None:
            food["salt"] = round(na * 2.5 / 1000, 3)  # mg sodium -> g salt
        foods.append(food)

    if len(foods) < 1000:
        sys.exit(f"Only {len(foods)} foods parsed; something's off")

    # Portion sizes ("1 slice", "1 cup"), so "2 slices of toast" can be logged in one go.
    measures = read_measures(links)
    with_measures = 0
    for food in foods:
        m = measures.get(food["k"])
        if m:
            food["m"] = m[:6]
            with_measures += 1
    print(f"Portion sizes for {with_measures} foods")
    with open(out, "w", encoding="utf-8") as fh:
        json.dump({
            "source": "AUSNUT 2023, Food Standards Australia New Zealand (CC BY-SA 3.0 AU based licence)",
            "foods": foods,
        }, fh, ensure_ascii=False, separators=(",", ":"))
    print(f"Wrote {len(foods)} foods to {out}")
    for f in foods[:5]:
        print("  e.g.", f)
    breads = [f for f in foods if "bread" in f["n"].lower()][:8]
    for f in breads:
        print("  bread:", f)


def read_measures(links):
    """{food key: [{"d": "1 slice", "g": 30.0}, ...]} from the 'Food measures' file, or {} if it can't be read."""
    candidates = [l for l in links if "measure" in l.lower()]
    if not candidates:
        print("!! No food measures file found; continuing without portion sizes")
        return {}
    url = candidates[0]
    print("Measures from:", url)
    try:
        wb = openpyxl.load_workbook(io.BytesIO(get(url)), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        print("!! couldn't read measures:", e)
        return {}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        for r, row in enumerate(rows[:15]):
            headers = [str(c) if c is not None else "" for c in row]
            key = pick(headers, "public food key", "food key", "food id")
            desc = pick(headers, "measure description", "measure|description", "description", "measure")
            grams = pick(headers, "weight|(g)", "weight", "gram", avoid=("description",))
            qty = pick(headers, "quantity", "number of")
            if key is None or desc is None or grams is None:
                continue
            print(f"Sheet '{ws.title}', header row {r + 1}:")
            for i, h in enumerate(headers):
                print(f"  [{i}] {h}")
            print("Measure columns:", {"key": headers[key], "desc": headers[desc], "grams": headers[grams], "qty": headers[qty] if qty is not None else None})
            out = {}
            for row in rows[r + 1:]:
                try:
                    k = str(row[key]).strip()
                    d = " ".join(str(row[desc]).split())
                    g = float(row[grams])
                except (TypeError, ValueError, IndexError):
                    continue
                if not k or not d or d == "None" or not (0 < g < 3000):
                    continue
                q = None
                if qty is not None and qty < len(row):
                    try:
                        q = float(row[qty])
                    except (TypeError, ValueError):
                        q = None
                label = d
                if q and q != 1 and not re.match(r"^\d", d):
                    label = f"{q:g} {d}"
                    g = g  # weight is for the stated quantity
                elif not re.match(r"^\d", d):
                    label = f"1 {d}"
                out.setdefault(k, []).append({"d": label, "g": round(g, 1)})
            sample = list(out.items())[:3]
            print("Measure examples:", sample)
            return out
    print("!! Couldn't find measure columns; continuing without portion sizes")
    return {}


if __name__ == "__main__":
    main()
